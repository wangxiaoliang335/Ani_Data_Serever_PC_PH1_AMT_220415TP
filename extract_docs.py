from docx import Document
from openpyxl import load_workbook
from pathlib import Path
import zipfile
import os

# 工程根目录
base_dir = r"E:\CSOT_PH1_AFT151_CAMERA\AMT 4# Data Server软件\Ani_Data_Serever_PC_PH1_AMT_220415TP"

docx_path = os.path.join(base_dir, "点灯检软件通信说明.docx")
xlsx_path = os.path.join(base_dir, "点灯数据库表说明.xlsx")

out_docx_txt = os.path.join(base_dir, "点灯检软件通信说明_导出.txt")
out_xlsx_txt = os.path.join(base_dir, "点灯数据库表说明_导出.txt")
out_docx_img_dir = os.path.join(base_dir, "点灯检软件通信说明_图片")


def export_docx_text():
    print("开始导出 Word 文本...")
    doc = Document(docx_path)
    lines = []

    # 段落
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)

    # 表格
    for ti, table in enumerate(doc.tables, start=1):
        lines.append(f"\n=== 表格 {ti} ===")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append("\t".join(cells))

    with open(out_docx_txt, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"已导出 Word 文本到: {out_docx_txt}")


def export_docx_images():
    print("开始导出 Word 图片...")
    img_dir = Path(out_docx_img_dir)
    img_dir.mkdir(exist_ok=True)

    # .docx 实际上是 zip 文件，图片都在 word/media 下面
    with zipfile.ZipFile(docx_path, "r") as zf:
        count = 0
        for name in zf.namelist():
            # 只处理真正的文件，跳过目录项
            if not name.startswith("word/media/"):
                continue
            if name.endswith("/"):  # 目录
                continue

            img_name = name.split("/")[-1]
            if not img_name:  # 安全起见再过滤一次
                continue

            data = zf.read(name)
            out_path = img_dir / img_name
            with open(out_path, "wb") as f:
                f.write(data)
            count += 1

    print(f"已从 Word 中导出 {count} 张图片到文件夹: {img_dir}，共 {count} 张。")


def export_xlsx():
    print("开始导出 Excel 文本...")
    wb = load_workbook(xlsx_path, data_only=True)
    lines = []
    for sheet in wb.worksheets:
        lines.append(f"\n=== 工作表: {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            vals = [str(v) if v is not None else "" for v in row]
            lines.append("\t".join(vals))

    with open(out_xlsx_txt, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"已导出 Excel 到: {out_xlsx_txt}")


if __name__ == "__main__":
    export_docx_text()
    export_docx_images()
    export_xlsx()
    print("全部导出完成。")